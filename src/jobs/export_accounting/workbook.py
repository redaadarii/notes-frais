"""Construction du classeur d'export comptable.

Module PUR : aucune dépendance à Spark ni à Databricks. Il reçoit des lignes déjà lues
et renvoie le contenu du fichier .xlsx ; il se teste donc en local avec pytest.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

AMOUNT_FORMAT = '#,##0.00 "€"'
DATE_FORMAT = "DD/MM/YYYY"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E293B")
TOTAL_FONT = Font(bold=True)


@dataclass(frozen=True)
class ExportLine:
    """Une ligne de dépense validée, enrichie de ce que la comptabilité attend."""

    report_id: str
    report_title: str
    employee_email: str
    employee_name: str
    cost_center: str
    expense_date: date
    gl_account: str
    category_label: str
    description: str
    amount: Decimal


def build_workbook(lines: list[ExportLine], period: str, batch_id: str) -> bytes:
    """Renvoie le contenu d'un classeur à trois onglets : écritures, synthèse, notes."""
    if not lines:
        raise ValueError("Aucune ligne à exporter")

    workbook = Workbook()
    _write_entries(workbook.active, lines)
    _write_summary(workbook.create_sheet(), lines, period, batch_id)
    _write_reports(workbook.create_sheet(), lines)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_entries(sheet: Worksheet, lines: list[ExportLine]) -> None:
    sheet.title = "Écritures"
    _write_header(sheet, ["Date", "Compte", "Centre de coût", "Collaborateur", "Note", "Catégorie", "Libellé", "Montant"])
    ordered = sorted(lines, key=lambda l: (l.gl_account, l.cost_center, l.expense_date))
    for line in ordered:
        sheet.append([
            line.expense_date,
            line.gl_account,
            line.cost_center,
            line.employee_name,
            line.report_title,
            line.category_label,
            line.description,
            line.amount,
        ])
    last_row = sheet.max_row
    _format_column(sheet, "A", DATE_FORMAT)
    _format_column(sheet, "H", AMOUNT_FORMAT)
    _append_total(sheet, label_column=7, amount_column=8, first_row=2, last_row=last_row)
    _finish(sheet, widths=[12, 10, 16, 24, 30, 24, 40, 14], last_row=last_row)


def _write_summary(sheet: Worksheet, lines: list[ExportLine], period: str, batch_id: str) -> None:
    sheet.title = "Synthèse"
    sheet.append(["Période", period])
    sheet.append(["Lot d'export", batch_id])

    totals: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
    for line in lines:
        totals[(line.gl_account, line.cost_center)] += line.amount

    header_row = 4  # une ligne vide sépare l'en-tête de la synthèse
    _write_header(sheet, ["Compte", "Centre de coût", "Montant"], row=header_row)
    for (account, cost_center), amount in sorted(totals.items()):
        sheet.append([account, cost_center, amount])
    last_row = sheet.max_row
    _format_column(sheet, "C", AMOUNT_FORMAT, first_row=header_row + 1)
    _append_total(sheet, label_column=2, amount_column=3, first_row=header_row + 1, last_row=last_row)
    _finish(sheet, widths=[14, 20, 16], header_row=header_row, last_row=last_row)


def _write_reports(sheet: Worksheet, lines: list[ExportLine]) -> None:
    sheet.title = "Notes"
    _write_header(sheet, ["Identifiant", "Titre", "Collaborateur", "E-mail", "Lignes", "Total"])

    reports: dict[str, list[ExportLine]] = defaultdict(list)
    for line in lines:
        reports[line.report_id].append(line)

    for report_id, report_lines in sorted(reports.items(), key=lambda item: item[1][0].employee_name):
        first = report_lines[0]
        sheet.append([
            report_id,
            first.report_title,
            first.employee_name,
            first.employee_email,
            len(report_lines),
            sum((l.amount for l in report_lines), Decimal("0")),
        ])
    _format_column(sheet, "F", AMOUNT_FORMAT)
    _finish(sheet, widths=[38, 30, 24, 32, 8, 14], last_row=sheet.max_row)


def _write_header(sheet: Worksheet, titles: list[str], row: int = 1) -> None:
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _format_column(sheet: Worksheet, letter: str, number_format: str, first_row: int = 2) -> None:
    for row in range(first_row, sheet.max_row + 1):
        sheet[f"{letter}{row}"].number_format = number_format


def _append_total(sheet: Worksheet, label_column: int, amount_column: int, first_row: int, last_row: int) -> None:
    """Total en FORMULE : le comptable qui modifie une ligne voit le total se mettre à jour."""
    total_row = last_row + 1
    letter = get_column_letter(amount_column)
    sheet.cell(row=total_row, column=label_column, value="Total").font = TOTAL_FONT
    total = sheet.cell(row=total_row, column=amount_column, value=f"=SUM({letter}{first_row}:{letter}{last_row})")
    total.font = TOTAL_FONT
    total.number_format = AMOUNT_FORMAT


def _finish(sheet: Worksheet, widths: list[int], last_row: int, header_row: int = 1) -> None:
    """Largeurs, en-tête figé et filtre automatique limité aux données (la ligne de total en est exclue)."""
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{last_row}"
