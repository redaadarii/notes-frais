from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from workbook import ExportLine, build_workbook


def line(**overrides) -> ExportLine:
    values = dict(
        report_id="r1",
        report_title="Salon Lyon",
        employee_email="reda@exemple.com",
        employee_name="Reda Adari",
        cost_center="VENTES",
        expense_date=date(2026, 9, 12),
        gl_account="6257",
        category_label="Repas",
        description="Déjeuner client",
        amount=Decimal("42.50"),
    )
    return ExportLine(**{**values, **overrides})


def open_workbook(content: bytes):
    return load_workbook(BytesIO(content))


def test_workbook_has_three_sheets():
    workbook = open_workbook(build_workbook([line()], "2026-09", "EXP-1"))
    assert workbook.sheetnames == ["Écritures", "Synthèse", "Notes"]


def test_entries_are_sorted_by_account_and_end_with_a_sum_formula():
    lines = [line(gl_account="6257"), line(gl_account="6241", amount=Decimal("10.00"))]
    sheet = open_workbook(build_workbook(lines, "2026-09", "EXP-1"))["Écritures"]

    assert [sheet["B2"].value, sheet["B3"].value] == ["6241", "6257"]
    assert sheet["G4"].value == "Total"
    assert sheet["H4"].value == "=SUM(H2:H3)"
    assert sheet.auto_filter.ref == "A1:H3"


def test_summary_groups_amounts_by_account_and_cost_center():
    lines = [
        line(amount=Decimal("10.00")),
        line(amount=Decimal("5.25")),
        line(cost_center="IT", amount=Decimal("1.00")),
    ]
    sheet = open_workbook(build_workbook(lines, "2026-09", "EXP-1"))["Synthèse"]

    assert sheet["B1"].value == "2026-09"
    assert sheet["B2"].value == "EXP-1"
    rows = [(r[0].value, r[1].value, r[2].value) for r in sheet.iter_rows(min_row=5, max_row=6)]
    assert rows == [("6257", "IT", 1), ("6257", "VENTES", 15.25)]


def test_reports_sheet_has_one_row_per_report():
    lines = [line(), line(), line(report_id="r2", employee_name="Omar Fassi")]
    sheet = open_workbook(build_workbook(lines, "2026-09", "EXP-1"))["Notes"]

    assert sheet.max_row == 3
    assert (sheet["A2"].value, sheet["E2"].value, sheet["F2"].value) == ("r2", 1, 42.5)
    assert (sheet["A3"].value, sheet["E3"].value, sheet["F3"].value) == ("r1", 2, 85)


def test_empty_export_is_refused():
    with pytest.raises(ValueError):
        build_workbook([], "2026-09", "EXP-1")
